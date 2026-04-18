import sys
import json
from pptx import Presentation
from pptx.util import Inches

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8")

def extract(filepath):
    prs = Presentation(filepath)
    slides = []
    for i, slide in enumerate(prs.slides):
        text_parts = []
        has_images = False

        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    line = para.text.strip()
                    if line:
                        text_parts.append(line)
            if shape.shape_type == 13:  # MSO_SHAPE_TYPE.PICTURE
                has_images = True

        notes_text = None
        if slide.has_notes_slide:
            notes_tf = slide.notes_slide.notes_text_frame
            notes_text = notes_tf.text.strip() if notes_tf else None
            if not notes_text:
                notes_text = None

        slides.append({
            "slideNumber": i + 1,
            "text": "\n".join(text_parts),
            "speakerNotes": notes_text,
            "hasImages": has_images
        })

    print(json.dumps({"slides": slides}, ensure_ascii=False))

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(json.dumps({"error": "No file path provided"}, ensure_ascii=False), file=sys.stderr)
        sys.exit(1)
    try:
        extract(sys.argv[1])
    except Exception as e:
        print(json.dumps({"error": str(e)}, ensure_ascii=False), file=sys.stderr)
        sys.exit(1)
